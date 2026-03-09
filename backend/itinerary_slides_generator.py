"""
Premium itinerary slide generator – VSL 360.

Design pillars:
  • Full-bleed imagery with cinematic gradient overlays
  • Generous white-space, editorial typography (mixed weights)
  • Gold / charcoal luxury palette with subtle warm neutrals
  • Layered depth: semi-transparent panels, thin accent lines,
    decorative geometric elements
  • Magazine-style alternating day layouts
"""

from datetime import datetime
import base64
import io
import mimetypes
import os
from typing import Dict, List, Optional, Tuple

import requests
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

from image_generator import auto_populate_images

# ── constants ────────────────────────────────────────────────────────────────
W, H = 7.5, 13.33  # portrait dims (inches) — mobile-friendly 9:16

LOCAL_IMAGE_DIR = os.path.abspath(
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "sample_itinerary", "images")
)

# ── colour palettes ─────────────────────────────────────────────────────────
THEMES = {
    "modern_luxe": {
        "bg":          RGBColor(250, 249, 247),   # warm off-white
        "ink":         RGBColor(28, 28, 30),       # near-black
        "sub":         RGBColor(72, 72, 74),       # dark grey
        "muted":       RGBColor(142, 142, 147),    # mid grey
        "accent":      RGBColor(183, 149, 93),     # warm gold
        "accent_dark": RGBColor(140, 109, 60),     # deep gold
        "accent_light":RGBColor(232, 218, 192),    # pale gold
        "line":        RGBColor(225, 221, 214),     # warm rule
        "panel":       RGBColor(255, 255, 255),
        "panel_alt":   RGBColor(245, 243, 240),    # tinted card
        "overlay_dark":RGBColor(18, 18, 20),
        "success":     RGBColor(52, 140, 98),
        "danger":      RGBColor(196, 69, 69),
        "serif":       "Georgia",
        "sans":        "Segoe UI",
    },
    "editorial_bold": {
        "bg":          RGBColor(248, 246, 240),
        "ink":         RGBColor(38, 34, 31),
        "sub":         RGBColor(82, 75, 68),
        "muted":       RGBColor(145, 136, 125),
        "accent":      RGBColor(183, 134, 62),
        "accent_dark": RGBColor(131, 86, 35),
        "accent_light":RGBColor(239, 224, 199),
        "line":        RGBColor(220, 212, 200),
        "panel":       RGBColor(255, 252, 248),
        "panel_alt":   RGBColor(247, 244, 238),
        "overlay_dark":RGBColor(24, 20, 16),
        "success":     RGBColor(74, 140, 91),
        "danger":      RGBColor(181, 74, 74),
        "serif":       "Georgia",
        "sans":        "Segoe UI",
    },
}

# ── image helpers ────────────────────────────────────────────────────────────

def _resolve_image_source_to_local(source: str) -> Tuple[Optional[str], bool]:
    if not source:
        return None, False
    normalized = source.strip()
    if normalized.startswith(("http://", "https://")):
        return None, False
    candidates = []
    if os.path.isabs(normalized):
        candidates.append(normalized)
    else:
        base = os.path.dirname(os.path.abspath(__file__))
        candidates.append(os.path.join(LOCAL_IMAGE_DIR, normalized))
        candidates.append(os.path.abspath(os.path.join(base, "..", normalized)))
        candidates.append(os.path.abspath(os.path.join(base, normalized)))
    for c in candidates:
        if os.path.isfile(c):
            return c, False
    return None, False


def _load_image_bytes(source: Optional[str]) -> Optional[bytes]:
    if not source:
        return None
    try:
        if source.startswith("data:image"):
            _, encoded = source.split(",", 1)
            return base64.b64decode(encoded)
        if source.startswith(("http://", "https://")):
            r = requests.get(source, timeout=20)
            r.raise_for_status()
            return r.content
        local, _ = _resolve_image_source_to_local(source)
        if local:
            with open(local, "rb") as f:
                return f.read()
    except Exception:
        return None
    return None


def _to_pptx_compatible_image_stream(image_bytes: bytes) -> io.BytesIO:
    with Image.open(io.BytesIO(image_bytes)) as img:
        if img.mode not in ("RGB", "RGBA"):
            img = img.convert("RGB")
        out = io.BytesIO()
        img.save(out, format="PNG")
        out.seek(0)
        return out


# ── low-level drawing primitives ─────────────────────────────────────────────

def _rect(slide, left, top, width, height, fill: RGBColor,
          line_rgb: Optional[RGBColor] = None, rounded=False, transparency: Optional[int] = None):
    """Draw a rectangle. *transparency* 0-100 (percent)."""
    from lxml import etree
    st = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(st, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if transparency is not None:
        # Inject alpha into the XML: spPr > solidFill > srgbClr > a:alpha
        sp_pr = shape._element.spPr
        solid_fill_elem = sp_pr.find(qn('a:solidFill'))
        if solid_fill_elem is not None:
            clr_elem = solid_fill_elem.find(qn('a:srgbClr'))
            if clr_elem is not None:
                alpha_elem = clr_elem.find(qn('a:alpha'))
                if alpha_elem is None:
                    alpha_elem = etree.SubElement(clr_elem, qn('a:alpha'))
                alpha_elem.set('val', str((100 - transparency) * 1000))
    if line_rgb:
        shape.line.color.rgb = line_rgb
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    return shape


def _thin_line(slide, left, top, width, color: RGBColor, thickness=0.02):
    """Ultra-thin horizontal rule."""
    return _rect(slide, left, top, width, thickness, color)


def _text(slide, text: str, left, top, width, height,
          size=24, color: Optional[RGBColor] = None, bold=False,
          align=PP_ALIGN.LEFT, font="Segoe UI", italic=False,
          spacing: Optional[float] = None, anchor=MSO_ANCHOR.TOP):
    """Add a text box with fine-grained control."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    if spacing is not None:
        p.space_after = Pt(spacing)
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    return box


def _multiline_text(slide, lines: List[str], left, top, width, height,
                    size=13, color: Optional[RGBColor] = None, font="Segoe UI",
                    line_spacing: float = 1.15, bold=False, align=PP_ALIGN.LEFT):
    """Text box with explicit line breaks and controlled leading."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(size * (line_spacing - 1))
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.name = font
            run.font.bold = bold
            if color:
                run.font.color.rgb = color
    return box


def _picture(slide, source: Optional[str], left, top, width, height) -> bool:
    """Place an image; returns True on success."""
    raw = _load_image_bytes(source)
    if not raw:
        return False
    try:
        stream = _to_pptx_compatible_image_stream(raw)
        slide.shapes.add_picture(stream, Inches(left), Inches(top), Inches(width), Inches(height))
        return True
    except Exception:
        return False


def _circle(slide, cx, cy, r, fill: RGBColor, line_rgb: Optional[RGBColor] = None):
    """Small decorative circle (centred on cx, cy)."""
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_rgb:
        shape.line.color.rgb = line_rgb
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


# ── theme selector ───────────────────────────────────────────────────────────

def _pick_theme(template_name: Optional[str]) -> dict:
    name = (template_name or "modern_luxe").lower()
    if "editorial" in name or "bold" in name:
        return THEMES["editorial_bold"]
    return THEMES["modern_luxe"]


def _safe_list(data, key: str):
    v = data.get(key, [])
    return v if isinstance(v, list) else []


# ═══════════════════════════════════════════════════════════════════════════════
#  S L I D E   B U I L D E R S  (portrait 7.5 × 13.33)
# ═══════════════════════════════════════════════════════════════════════════════

def _add_cover_slide(prs, trip_data, t):
    """Full-bleed portrait cover — centred branding, duration pill, website."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Full-bleed background image
    src = trip_data.get("cover_image_path") or trip_data.get("cover_image_url")
    if not _picture(slide, src, 0, 0, W, H):
        _rect(slide, 0, 0, W, H, t["overlay_dark"])

    # Subtle dark overlay for text legibility
    _rect(slide, 0, 0, W, H, RGBColor(0, 0, 0), transparency=60)

    # ── Brand block — top centre ──
    _text(slide, "VSL 360", 0, 1.2, W, 0.55, 30, RGBColor(255, 255, 255), True,
          PP_ALIGN.CENTER, font=t["serif"])
    _text(slide, "E X P L O R E   E V E R Y   A N G L E", 0, 1.75, W, 0.3, 8,
          RGBColor(210, 210, 210), False, PP_ALIGN.CENTER, font=t["sans"])

    # Trip name — large elegant centred italic
    trip_name = trip_data.get("trip_name", "Travel Itinerary")
    _text(slide, trip_name, 0.4, 5.0, W - 0.8, 1.6, 52, RGBColor(255, 255, 255), True,
          PP_ALIGN.CENTER, font=t["serif"], italic=True)

    # Duration pill — rounded outline badge
    days = trip_data.get("total_days", 0)
    nights = trip_data.get("total_nights", 0)
    pill_text = f"{days} DAYS  |  {nights} NIGHTS"
    pill_w = 3.8
    pill_x = (W - pill_w) / 2
    _rect(slide, pill_x, 7.0, pill_w, 0.55, RGBColor(0, 0, 0),
          RGBColor(255, 255, 255), rounded=True, transparency=80)
    _text(slide, pill_text, pill_x, 7.0, pill_w, 0.55, 12,
          RGBColor(255, 255, 255), True, PP_ALIGN.CENTER, font=t["sans"],
          anchor=MSO_ANCHOR.MIDDLE)

    # Customer name
    cust = trip_data.get("customer_name")
    if cust:
        _text(slide, f"Prepared for {cust}", 0, 11.5, W, 0.35, 11,
              RGBColor(200, 200, 200), False, PP_ALIGN.CENTER, font=t["sans"],
              italic=True)

    # Website at bottom
    _text(slide, "www.visitsrilanka360.com", 0, 12.2, W, 0.35, 11,
          RGBColor(220, 220, 220), False, PP_ALIGN.CENTER, font=t["sans"],
          italic=True)


def _add_overview_slide(prs, trip_data, t):
    """Portrait overview — stacked layout with destination card."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, W, H, t["bg"])

    # Gold accent bar at top
    _rect(slide, 0, 0, W, 0.06, t["accent"])

    # Section label
    _text(slide, "OVERVIEW", 0.6, 0.55, 4, 0.3, 10, t["muted"], True, font=t["sans"])
    _thin_line(slide, 0.6, 0.88, 1.2, t["accent"])

    # Title
    _text(slide, "Your Journey", 0.6, 1.1, 6, 0.65, 34, t["ink"], True, font=t["serif"])

    # Overview text
    overview = trip_data.get("overview_text", "") or \
        "A carefully curated journey through breathtaking landscapes and rich cultural heritage."
    _text(slide, overview, 0.6, 2.0, W - 1.2, 2.5, 14, t["sub"], False, font=t["sans"])

    # Dates
    start = trip_data.get("start_date", "")
    end = trip_data.get("end_date", "")
    if start and end:
        _text(slide, f"{start}  —  {end}", 0.6, 4.4, 5, 0.3, 11, t["muted"], False,
              font=t["sans"], italic=True)

    # Destinations card — full-width below overview text
    card_top = 5.1
    card_w = W - 1.2
    _rect(slide, 0.6, card_top, card_w, 7.5, t["panel"], t["line"], rounded=True)
    _rect(slide, 0.6, card_top, card_w, 0.06, t["accent"])

    _text(slide, "DESTINATIONS", 1.0, card_top + 0.35, 4, 0.3, 10, t["accent_dark"], True,
          font=t["sans"])
    _thin_line(slide, 1.0, card_top + 0.7, 2.0, t["line"])

    y = card_top + 1.0
    for item in _safe_list(trip_data, "location_summary")[:10]:
        city = item.get("city", "")
        nights = item.get("nights", 0)
        _circle(slide, 1.15, y + 0.13, 0.05, t["accent"])
        _text(slide, city, 1.4, y, 3.5, 0.3, 14, t["ink"], True, font=t["sans"])
        _text(slide, f"{nights} night{'s' if nights != 1 else ''}", 5.2, y, 1.5, 0.3,
              11, t["muted"], False, PP_ALIGN.RIGHT, font=t["sans"])
        y += 0.48

    # Total
    _thin_line(slide, 1.0, y + 0.15, card_w - 0.8, t["line"])
    _text(slide, f"{trip_data.get('total_nights', 0)} Nights  |  {trip_data.get('total_days', 0)} Days",
          1.0, y + 0.35, card_w - 0.8, 0.35, 13, t["accent_dark"], True,
          PP_ALIGN.CENTER, font=t["sans"])


def _add_route_slide(prs, trip_data, t):
    """Portrait vertical route timeline."""
    days = _safe_list(trip_data, "days")
    if not days:
        return

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, W, H, t["panel"])
    _rect(slide, 0, 0, W, 0.06, t["accent"])

    _text(slide, "ROUTE", 0.6, 0.55, 3, 0.3, 10, t["muted"], True, font=t["sans"])
    _thin_line(slide, 0.6, 0.88, 1.0, t["accent"])
    _text(slide, "At a Glance", 0.6, 1.1, 5, 0.55, 34, t["ink"], True, font=t["serif"])

    # Extract unique consecutive cities
    stops = []
    for d in days:
        city = (d.get("city") or "").strip()
        if city and (not stops or stops[-1] != city):
            stops.append(city)
    if not stops:
        stops = ["Start", "Destination"]
    max_n = 10
    if len(stops) > max_n:
        stops = stops[:max_n - 1] + [stops[-1]]

    # Vertical timeline
    rail_x = 1.4
    rail_top = 2.2
    rail_bottom = 11.5
    rail_h = rail_bottom - rail_top

    # Vertical connector line
    _rect(slide, rail_x - 0.015, rail_top, 0.03, rail_h, t["line"])

    spacing = rail_h / max(1, len(stops) - 1)
    for i, city in enumerate(stops):
        cy = rail_top + i * spacing
        # Outer ring
        _circle(slide, rail_x, cy, 0.14, t["panel"], t["accent"])
        # Inner dot
        _circle(slide, rail_x, cy, 0.07, t["accent"])
        # City label to the right
        _text(slide, city, 1.9, cy - 0.15, 4.0, 0.3, 14, t["ink"], True, font=t["sans"])
        # Day number
        day_idx = i + 1
        if day_idx <= len(days):
            _text(slide, f"Day {day_idx}", 1.9, cy + 0.15, 2.0, 0.25, 9, t["muted"],
                  False, font=t["sans"])

    # Summary
    _rect(slide, 0.6, 12.0, W - 1.2, 0.8, t["panel_alt"], t["line"], rounded=True)
    _text(slide, f"{trip_data.get('total_days', 0)} days  •  {trip_data.get('total_nights', 0)} nights  •  Curated experiences",
          0.6, 12.15, W - 1.2, 0.5, 12, t["muted"], False, PP_ALIGN.CENTER, font=t["sans"])


def _add_day_pair_slide(prs, day_a, idx_a, day_b, idx_b, t):
    """Two days stacked on one portrait slide — alternating image/card."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Warm beige background (visible at edges and between halves)
    _rect(slide, 0, 0, W, H, t["accent_light"])

    half_h = H / 2  # ~6.665
    gap = 0.12

    # Top day: card left, image right
    _draw_day_half(slide, day_a, idx_a, t, y_offset=0, half_h=half_h - gap / 2,
                   image_right=True)
    # Bottom day: image left, card right
    _draw_day_half(slide, day_b, idx_b, t, y_offset=half_h + gap / 2,
                   half_h=half_h - gap / 2, image_right=False)


def _add_day_single_slide(prs, day, idx, t, image_right=True):
    """Single day on a full portrait slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, W, H, t["accent_light"])
    _draw_day_full(slide, day, idx, t, image_right=image_right)


def _draw_day_full(slide, day, idx, t, image_right=True):
    """Single day filling the entire portrait slide — generous layout."""
    m = 0.2  # outer margin

    # Image: fills ~55% width, full height minus margins
    img_w = 4.2
    img_h = H - 2 * m
    if image_right:
        img_x = W - img_w - m
        card_x = m
    else:
        img_x = m
        card_x = W - 5.2 - m

    # Place image (drawn first, card overlaps it)
    img_src = day.get("image_path") or day.get("image_url")
    placed = _picture(slide, img_src, img_x, m, img_w, img_h) if img_src else False
    if not placed:
        _rect(slide, img_x, m, img_w, img_h, t["panel_alt"])

    # White rounded card — 70% width, overlaps image
    card_w = 5.2
    card_top = m + 0.3
    card_h = H - 2 * m - 0.6
    _rect(slide, card_x, card_top, card_w, card_h, t["panel"], rounded=True)

    # Content
    pad = 0.5
    cx = card_x + pad
    cw = card_w - 2 * pad

    day_num = day.get("day", idx + 1)
    _text(slide, f"Day {day_num}", cx, card_top + 0.5, cw, 0.7, 42,
          t["ink"], True, font=t["serif"], italic=True)

    title = day.get("title", "Untitled Day")
    _text(slide, title, cx, card_top + 1.3, cw, 0.5, 22,
          t["accent_dark"], True, font=t["serif"])

    activities = day.get("activities", []) or []
    ay = card_top + 2.2
    for act in activities[:10]:
        desc = (act.get("description") or "").strip() or "Activity"
        time_s = (act.get("time") or "").strip()
        bullet_text = f"{desc} ({time_s})" if time_s else desc
        _circle(slide, cx + 0.12, ay + 0.14, 0.05, t["ink"])
        _text(slide, bullet_text, cx + 0.35, ay, cw - 0.35, 0.35,
              14, t["sub"], False, font=t["sans"])
        ay += 0.48

    optional = day.get("optional")
    if optional:
        ay += 0.2
        _text(slide, f"Optional : {optional}", cx, ay, cw, 0.3, 13,
              t["sub"], False, font=t["sans"])

    # Bottom stats
    bottom_y = card_top + card_h - 1.1
    _thin_line(slide, cx, bottom_y, cw, t["line"])

    tt = day.get("travel_time")
    dist = day.get("distance")
    left_lines = []
    if tt and tt != "N/A":
        left_lines.append(f"Travel Time: {tt}")
    if dist and dist != "N/A":
        left_lines.append(f"Distance: {dist}")
    if left_lines:
        _multiline_text(slide, left_lines, cx, bottom_y + 0.15, cw / 2, 0.7,
                        12, t["accent_dark"], font=t["sans"], bold=True,
                        line_spacing=1.4)

    overnight = day.get("overnight_city") or day.get("city")
    if overnight:
        _text(slide, "Overnight stay", cx + cw / 2, bottom_y + 0.15, cw / 2, 0.25, 11,
              t["sub"], False, PP_ALIGN.RIGHT, font=t["sans"], italic=True)
        _text(slide, overnight, cx + cw / 2, bottom_y + 0.45, cw / 2, 0.4, 18,
              t["ink"], True, PP_ALIGN.RIGHT, font=t["serif"], italic=True)


def _draw_day_half(slide, day, idx, t, y_offset, half_h, image_right=True):
    """Draw one day within a vertical half of a portrait slide.

    Image fills one side edge-to-edge, white rounded card overlaps it
    from the other side — matching the reference design.
    """
    m = 0.2  # outer margin

    # Image fills ~55% of width, full half height minus small margin
    img_w = 4.0
    img_h = half_h - 2 * m
    if image_right:
        img_x = W - img_w - m
        card_x = m
    else:
        img_x = m
        card_x = W - 4.8 - m

    img_y = y_offset + m

    # Place image first (card will overlap)
    img_src = day.get("image_path") or day.get("image_url")
    placed = _picture(slide, img_src, img_x, img_y, img_w, img_h) if img_src else False
    if not placed:
        _rect(slide, img_x, img_y, img_w, img_h, t["panel_alt"])

    # White rounded card — overlaps image by ~1.5"
    card_w = 4.8
    card_top = y_offset + m + 0.15
    card_h = half_h - 2 * m - 0.3
    _rect(slide, card_x, card_top, card_w, card_h, t["panel"], rounded=True)

    # ── Card content ──
    pad = 0.45
    cx = card_x + pad
    cw = card_w - 2 * pad

    # "Day X" — bold italic serif, large
    day_num = day.get("day", idx + 1)
    _text(slide, f"Day {day_num}", cx, card_top + 0.3, cw, 0.55, 32,
          t["ink"], True, font=t["serif"], italic=True)

    # Title / subtitle in gold bold
    title = day.get("title", "Untitled Day")
    city = day.get("city", "")
    display_title = title
    if city and city.lower() not in title.lower():
        display_title = f"{title}"  # city shown via overnight at bottom
    _text(slide, display_title, cx, card_top + 0.9, cw, 0.4, 16,
          t["accent_dark"], True, font=t["serif"])

    # Activities — bullet list
    activities = day.get("activities", []) or []
    ay = card_top + 1.5
    max_acts = 6
    for act in activities[:max_acts]:
        desc = (act.get("description") or "").strip() or "Activity"
        time_s = (act.get("time") or "").strip()
        bullet_text = f"{desc} ({time_s})" if time_s else desc
        _circle(slide, cx + 0.1, ay + 0.12, 0.045, t["ink"])
        _text(slide, bullet_text, cx + 0.3, ay, cw - 0.3, 0.3,
              12, t["sub"], False, font=t["sans"])
        ay += 0.38

    # Optional note
    optional = day.get("optional")
    if optional:
        ay += 0.1
        _text(slide, f"Optional : {optional}", cx, ay, cw, 0.28, 11,
              t["sub"], False, font=t["sans"])

    # ── Bottom stats — two groups ──
    bottom_y = card_top + card_h - 0.8
    _thin_line(slide, cx, bottom_y - 0.05, cw, t["line"])

    # Left: Travel Time & Distance (gold bold)
    tt = day.get("travel_time")
    dist = day.get("distance")
    left_lines = []
    if tt and tt != "N/A":
        left_lines.append(f"Travel Time: {tt}")
    if dist and dist != "N/A":
        left_lines.append(f"Distance: {dist}")
    if left_lines:
        _multiline_text(slide, left_lines, cx, bottom_y + 0.05, cw / 2, 0.65,
                        10, t["accent_dark"], font=t["sans"], bold=True,
                        line_spacing=1.3)

    # Right: Overnight stay
    overnight = day.get("overnight_city") or day.get("city")
    if overnight:
        _text(slide, "Overnight stay", cx + cw / 2, bottom_y + 0.05, cw / 2, 0.22, 10,
              t["sub"], False, PP_ALIGN.RIGHT, font=t["sans"], italic=True)
        _text(slide, overnight, cx + cw / 2, bottom_y + 0.32, cw / 2, 0.35, 14,
              t["ink"], True, PP_ALIGN.RIGHT, font=t["serif"], italic=True)


def _add_hotels_slide(prs, trip_data, t):
    """Portrait accommodation list — clean card rows."""
    hotels = _safe_list(trip_data, "hotels")
    if not hotels:
        return

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, W, H, t["bg"])
    _rect(slide, 0, 0, W, 0.06, t["accent"])

    _text(slide, "ACCOMMODATION", 0.6, 0.55, 4, 0.3, 10, t["muted"], True, font=t["sans"])
    _thin_line(slide, 0.6, 0.88, 1.5, t["accent"])
    _text(slide, "Where You'll Stay", 0.6, 1.1, 6, 0.55, 34, t["ink"], True, font=t["serif"])

    y = 2.1
    card_w = W - 1.2
    for i, hotel in enumerate(hotels[:8]):
        row_bg = t["panel_alt"] if i % 2 == 0 else t["panel"]
        row_h = 1.15
        _rect(slide, 0.6, y, card_w, row_h, row_bg, t["line"], rounded=True)

        # Gold accent dot
        _circle(slide, 0.95, y + 0.3, 0.055, t["accent"])

        # Hotel name
        _text(slide, hotel.get("name", "Hotel"), 1.2, y + 0.15, card_w - 1.0, 0.35, 14,
              t["ink"], True, font=t["sans"])
        # Location + room type on second line
        details = []
        city = hotel.get("city", "")
        room = hotel.get("room_type", "")
        if city:
            details.append(city)
        if room:
            details.append(room)
        if details:
            _text(slide, "  •  ".join(details), 1.2, y + 0.52, card_w - 2.0, 0.3, 11,
                  t["muted"], False, font=t["sans"])

        # Nights badge on right
        night_val = hotel.get("night", "")
        if night_val:
            _text(slide, night_val, card_w - 0.8, y + 0.15, 1.6, 0.35, 13,
                  t["accent_dark"], True, PP_ALIGN.RIGHT, font=t["sans"])

        y += row_h + 0.12


def _add_inclusions_slide(prs, trip_data, t):
    """Portrait inclusions — stacked include/exclude cards."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, W, H, t["panel"])
    _rect(slide, 0, 0, W, 0.06, t["accent"])

    _text(slide, "PACKAGE DETAILS", 0.6, 0.55, 4, 0.3, 10, t["muted"], True, font=t["sans"])
    _thin_line(slide, 0.6, 0.88, 1.5, t["accent"])
    _text(slide, "What's Included", 0.6, 1.1, 6, 0.55, 34, t["ink"], True, font=t["serif"])

    card_w = W - 1.2

    # ── Included card ──
    inc_top = 2.0
    includes = _safe_list(trip_data, "tour_includes")
    inc_items = min(len(includes), 12)
    inc_h = 0.9 + inc_items * 0.32

    _rect(slide, 0.6, inc_top, card_w, inc_h, t["panel_alt"], t["line"], rounded=True)
    _thin_line(slide, 0.6, inc_top + 0.005, card_w, t["success"], thickness=0.04)
    _text(slide, "INCLUDED", 0.95, inc_top + 0.2, 4, 0.3, 10, t["success"], True, font=t["sans"])
    _thin_line(slide, 0.95, inc_top + 0.55, 1.5, t["line"])

    y = inc_top + 0.7
    for item in includes[:12]:
        _circle(slide, 1.1, y + 0.1, 0.04, t["success"])
        _text(slide, item, 1.3, y, card_w - 1.0, 0.25, 11, t["sub"], False, font=t["sans"])
        y += 0.32

    # ── Excluded card ──
    exc_top = inc_top + inc_h + 0.3
    excludes = _safe_list(trip_data, "tour_excludes")
    exc_items = min(len(excludes), 12)
    exc_h = 0.9 + exc_items * 0.32

    _rect(slide, 0.6, exc_top, card_w, exc_h, t["panel_alt"], t["line"], rounded=True)
    _thin_line(slide, 0.6, exc_top + 0.005, card_w, t["danger"], thickness=0.04)
    _text(slide, "NOT INCLUDED", 0.95, exc_top + 0.2, 4, 0.3, 10, t["danger"], True, font=t["sans"])
    _thin_line(slide, 0.95, exc_top + 0.55, 1.5, t["line"])

    y = exc_top + 0.7
    for item in excludes[:12]:
        _circle(slide, 1.1, y + 0.1, 0.04, t["danger"])
        _text(slide, item, 1.3, y, card_w - 1.0, 0.25, 11, t["sub"], False, font=t["sans"])
        y += 0.32

    # Cost badge
    cost = trip_data.get("cost_per_person")
    if cost:
        badge_w = 4.5
        badge_x = (W - badge_w) / 2
        badge_y = exc_top + exc_h + 0.4
        _rect(slide, badge_x, badge_y, badge_w, 0.45, t["accent"], rounded=True)
        _text(slide, f"Per Person:  {cost}", badge_x, badge_y + 0.02, badge_w, 0.4, 13,
              RGBColor(255, 255, 255), True, PP_ALIGN.CENTER, font=t["sans"])


def _add_closing_slide(prs, trip_data, t):
    """Portrait closing — full-bleed image with centred thank-you."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    src = trip_data.get("cover_image_path") or trip_data.get("cover_image_url")
    if not _picture(slide, src, 0, 0, W, H):
        _rect(slide, 0, 0, W, H, t["overlay_dark"])

    _rect(slide, 0, 0, W, H, RGBColor(0, 0, 0), transparency=45)

    # Centred content
    _thin_line(slide, (W - 2.5) / 2, 4.8, 2.5, t["accent"])

    _text(slide, "Thank You", 0, 5.2, W, 0.9, 48, RGBColor(255, 255, 255), True,
          PP_ALIGN.CENTER, font=t["serif"])

    _text(slide, trip_data.get("trip_name", ""),
          0, 6.3, W, 0.5, 18, t["accent_light"], False, PP_ALIGN.CENTER, font=t["sans"],
          italic=True)

    _thin_line(slide, (W - 2.5) / 2, 7.1, 2.5, t["accent"])

    _text(slide, "Your journey begins here.", 0, 7.5, W, 0.45, 14,
          RGBColor(200, 200, 200), False, PP_ALIGN.CENTER, font=t["sans"])

    # Brand
    _text(slide, "VSL 360", 0, 11.5, W, 0.4, 14, t["accent"], True, PP_ALIGN.CENTER,
          font=t["serif"])
    _text(slide, "www.visitsrilanka360.com", 0, 12.0, W, 0.35, 10,
          RGBColor(200, 200, 200), False, PP_ALIGN.CENTER, font=t["sans"])


# ═══════════════════════════════════════════════════════════════════════════════
#  P U B L I C   A P I
# ═══════════════════════════════════════════════════════════════════════════════

def generate_vsl360_itinerary_slides(trip_data: Dict, output_path: Optional[str] = None) -> str:
    if not output_path:
        trip_name = trip_data.get("trip_name", "itinerary").replace(" ", "_").lower()
        out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "itineraries")
        output_path = os.path.join(out_dir, f"{trip_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    trip_data = auto_populate_images(dict(trip_data))
    theme = _pick_theme(trip_data.get("template_name"))

    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    _add_cover_slide(prs, trip_data, theme)
    _add_overview_slide(prs, trip_data, theme)
    _add_route_slide(prs, trip_data, theme)

    # Day slides — pair two days per portrait slide (like the reference style)
    days = _safe_list(trip_data, "days")
    i = 0
    while i < len(days):
        if i + 1 < len(days):
            _add_day_pair_slide(prs, days[i], i, days[i + 1], i + 1, theme)
            i += 2
        else:
            _add_day_single_slide(prs, days[i], i, theme,
                                  image_right=(i % 2 == 0))
            i += 1

    _add_hotels_slide(prs, trip_data, theme)
    _add_inclusions_slide(prs, trip_data, theme)
    _add_closing_slide(prs, trip_data, theme)

    prs.save(output_path)
    return output_path
